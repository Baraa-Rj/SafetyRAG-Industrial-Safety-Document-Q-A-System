from __future__ import annotations

import re
import subprocess
from dataclasses import dataclass
from pathlib import Path

from .config import Source

TEXT_SUFFIXES = {".md", ".txt", ".py", ".json", ".yaml", ".yml", ".cfg", ".ini", ".rst"}
PDFTOTEXT_TIMEOUT = 120


@dataclass(frozen=True)
class RawDoc:
    path: Path
    label: str
    kind: str
    text: str
    fingerprint: str


@dataclass(frozen=True)
class LoadFailure:
    path: Path
    reason: str


def _fingerprint(path: Path) -> str:
    stat = path.stat()
    return f"{int(stat.st_mtime)}:{stat.st_size}"


def _read_pdf(path: Path) -> str:
    result = subprocess.run(
        ["pdftotext", "-layout", "-nopgbrk", str(path), "-"],
        capture_output=True,
        timeout=PDFTOTEXT_TIMEOUT,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.decode("utf-8", "replace").strip() or "pdftotext failed")
    return result.stdout.decode("utf-8", "replace")


def _read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


_TEX_DROP_ENVS = ("figure", "table", "tabular", "lstlisting", "verbatim", "tikzpicture")
_TEX_HEADINGS = (
    (r"\\chapter\*?\{([^}]*)\}", "# "),
    (r"\\section\*?\{([^}]*)\}", "## "),
    (r"\\subsection\*?\{([^}]*)\}", "### "),
    (r"\\subsubsection\*?\{([^}]*)\}", "#### "),
)


def _read_tex(path: Path) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")
    text = re.sub(r"(?<!\\)%.*", "", text)
    for env in _TEX_DROP_ENVS:
        text = re.sub(rf"\\begin\{{{env}\*?\}}.*?\\end\{{{env}\*?\}}", " ", text, flags=re.S)
    for pattern, prefix in _TEX_HEADINGS:
        text = re.sub(pattern, lambda m, p=prefix: f"\n\n{p}{m.group(1)}\n", text)
    text = re.sub(r"\\(?:label|ref|cite[a-z]*|index)\{[^}]*\}", " ", text)
    text = re.sub(r"\\(?:textbf|textit|emph|texttt|underline)\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\begin\{itemize\}|\\end\{itemize\}|\\begin\{enumerate\}|\\end\{enumerate\}", "", text)
    text = re.sub(r"\\item\s*", "- ", text)
    text = re.sub(r"\\[a-zA-Z]+\*?(\[[^\]]*\])?(\{[^}]*\})?", " ", text)
    text = text.replace("~", " ").replace("\\&", "&")
    return re.sub(r"\n{3,}", "\n\n", text)


def load_file(path: Path, label: str, kind: str) -> RawDoc:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _read_pdf(path)
    elif suffix == ".docx":
        text = _read_docx(path)
    elif suffix == ".tex":
        text = _read_tex(path)
    elif suffix in TEXT_SUFFIXES:
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"unsupported file type: {suffix or path.name}")
    return RawDoc(
        path=path,
        label=label,
        kind=kind,
        text=text,
        fingerprint=_fingerprint(path),
    )


def load_sources(sources: list[Source]) -> tuple[list[RawDoc], list[LoadFailure]]:
    docs: list[RawDoc] = []
    failures: list[LoadFailure] = []
    for source in sources:
        files = source.iter_files()
        if not files and not source.root.is_dir():
            failures.append(LoadFailure(source.root, f"source root missing ({source.label})"))
            continue
        for path in files:
            try:
                doc = load_file(path, source.label, source.kind)
            except Exception as exc:
                failures.append(LoadFailure(path, f"{type(exc).__name__}: {exc}"))
                continue
            if doc.text.strip():
                docs.append(doc)
            else:
                failures.append(LoadFailure(path, "no extractable text"))
    return docs, failures
